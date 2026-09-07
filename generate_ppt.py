from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x0F, 0x2A, 0x43)
BLUE = RGBColor(0x1F, 0x5C, 0x8A)
TEAL = RGBColor(0x1A, 0x7A, 0x6E)
GRAY = RGBColor(0x4A, 0x4A, 0x4A)
LIGHT_GRAY = RGBColor(0xF2, 0xF3, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SS = "dashboard/screenshots/"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

def add_slide():
    return prs.slides.add_slide(blank)

def set_bg(slide, color=WHITE):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color

def add_rect(slide, x, y, w, h, color, line=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if not line:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = color
    shp.shadow.inherit = False
    return shp

def add_text(slide, x, y, w, h, text, size=18, color=GRAY, bold=False,
             align=PP_ALIGN.LEFT, font="Calibri", anchor=None, line_spacing=1.0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    return tb

def add_bullets(slide, x, y, w, h, items, size=16, color=GRAY, bold_lead=True,
                 space_after=10, bullet_color=BLUE):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space_after)
        p.line_spacing = 1.15
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run()
            r1.text = "▪  " + lead
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = bullet_color
            r1.font.name = "Calibri"
            if rest:
                r2 = p.add_run()
                r2.text = " " + rest
                r2.font.size = Pt(size)
                r2.font.bold = False
                r2.font.color.rgb = color
                r2.font.name = "Calibri"
        else:
            r1 = p.add_run()
            r1.text = "▪  " + item
            r1.font.size = Pt(size)
            r1.font.color.rgb = color
            r1.font.name = "Calibri"
    return tb

def add_header(slide, kicker, title):
    add_rect(slide, 0, 0, prs.slide_width, Inches(1.15), NAVY)
    add_text(slide, Inches(0.55), Inches(0.12), Inches(8), Inches(0.4),
              kicker.upper(), size=13, color=RGBColor(0xB9, 0xD3, 0xE8), bold=True)
    add_text(slide, Inches(0.5), Inches(0.42), Inches(11), Inches(0.65),
              title, size=30, color=WHITE, bold=True)

def add_picture_framed(slide, path, x, y, w, h, caption=None):
    pic = slide.shapes.add_picture(path, x, y, height=h)
    if pic.width > w:
        ratio = w / pic.width
        pic.width = int(pic.width * ratio)
        pic.height = int(pic.height * ratio)
    pic.left = int(x + (w - pic.width) / 2)
    pic.top = int(y)
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, pic.left, pic.top, pic.width, pic.height)
    frame.fill.background()
    frame.line.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    frame.line.width = Pt(1)
    frame.shadow.inherit = False
    if caption:
        add_text(slide, pic.left, pic.top + pic.height + Inches(0.05), pic.width, Inches(0.35),
                  caption, size=12, color=GRAY, align=PP_ALIGN.CENTER)
    return pic

# ---------- Slide 1: Title ----------
s = add_slide()
set_bg(s, NAVY)
add_rect(s, 0, Inches(4.6), prs.slide_width, Inches(0.06), TEAL)
add_text(s, Inches(1), Inches(2.7), Inches(11.3), Inches(1.3),
          "Relief Dashboard", size=54, color=WHITE, bold=True)
add_text(s, Inches(1), Inches(3.85), Inches(11.3), Inches(0.7),
          "Real time disaster relief logistics coordination", size=22, color=RGBColor(0xB9, 0xD3, 0xE8))
add_text(s, Inches(1), Inches(6.6), Inches(11.3), Inches(0.5),
          "Flutter and Firebase application connecting field volunteers and coordinators", size=14, color=RGBColor(0x8C, 0xA8, 0xC2))

# ---------- Slide 2: The Problem ----------
s = add_slide()
set_bg(s)
add_header(s, "The Problem", "Relief coordination is fragmented and manual")
add_text(s, Inches(0.55), Inches(1.45), Inches(11.8), Inches(0.9),
          "When a disaster strikes, such as a flood, earthquake, or humanitarian crisis, relief "
          "operations are often coordinated through phone calls, group chats, paper forms, and "
          "spreadsheets passed between field workers and coordinators.",
          size=16, color=GRAY, line_spacing=1.2)
add_bullets(s, Inches(0.55), Inches(2.6), Inches(11.8), Inches(4),
    [
        ("No real time visibility.", "Coordinators cannot see what supplies exist, where, or in what quantity."),
        ("Slow, error prone request handling.", "Volunteers cannot easily report community needs, and coordinators cannot see requests as they arrive."),
        ("Misallocated resources.", "Stock is sent to the wrong location, or critical needs go unmet while other areas are over supplied."),
        ("No prioritization.", "Without urgency tracking, life threatening requests can sit in the same queue as low priority ones."),
    ], size=17)

# ---------- Slide 3: Who It Affects ----------
s = add_slide()
set_bg(s)
add_header(s, "The Problem", "Who it affects")
cards = [
    ("Field Volunteers", "Need a fast way to report needs and locations without technical overhead."),
    ("Coordinators and Admins", "Responsible for inventory, allocation, and making sure aid reaches the right place in time."),
    ("Affected Communities", "Wait longer or receive less effective aid when coordination breaks down."),
]
card_w = Inches(3.7)
gap = Inches(0.35)
start_x = Inches(0.55)
for i, (title, desc) in enumerate(cards):
    x = start_x + i * (card_w + gap)
    y = Inches(2.0)
    h = Inches(3.6)
    add_rect(s, x, y, card_w, h, LIGHT_GRAY)
    add_rect(s, x, y, card_w, Inches(0.08), TEAL)
    add_text(s, x + Inches(0.25), y + Inches(0.4), card_w - Inches(0.5), Inches(0.9),
              title, size=19, color=NAVY, bold=True, line_spacing=1.1)
    add_text(s, x + Inches(0.25), y + Inches(1.35), card_w - Inches(0.5), Inches(2.0),
              desc, size=14, color=GRAY, line_spacing=1.25)

# ---------- Slide 4: The Solution ----------
s = add_slide()
set_bg(s)
add_header(s, "The Solution", "A shared, live system of record")
add_text(s, Inches(0.55), Inches(1.45), Inches(11.8), Inches(0.7),
          "Relief Dashboard is a role based, real time logistics application that connects field "
          "volunteers directly to coordinators.", size=16, color=GRAY, line_spacing=1.2)
add_bullets(s, Inches(0.55), Inches(2.3), Inches(6.4), Inches(4),
    [
        ("Volunteers submit requests", "with location, photo evidence, urgency level, and needed items directly from the field."),
        ("Coordinators act in real time,", "checking inventory across warehouses and allocating stock with one tap."),
        ("Both roles share a live map view", "of requests and warehouses for full geographic context."),
    ], size=16)
add_picture_framed(s, SS + "volunteer-home.png", Inches(7.4), Inches(1.55), Inches(5.3), Inches(5.5),
                    "New relief request form")

# ---------- Slide 5: Who It Serves ----------
s = add_slide()
set_bg(s)
add_header(s, "The Solution", "Who it serves")
add_text(s, Inches(0.55), Inches(1.4), Inches(5.7), Inches(0.9),
          "Field Volunteer role", size=20, color=NAVY, bold=True)
add_text(s, Inches(0.55), Inches(1.9), Inches(5.7), Inches(0.9),
          "A lightweight, mobile friendly flow for reporting needs from the field.",
          size=15, color=GRAY, line_spacing=1.2)
add_picture_framed(s, SS + "volunteer-map.png", Inches(0.55), Inches(2.5), Inches(5.7), Inches(4.4),
                    "Volunteer map view")

add_text(s, Inches(6.9), Inches(1.4), Inches(5.9), Inches(0.9),
          "Coordinator / Admin role", size=20, color=NAVY, bold=True)
add_text(s, Inches(6.9), Inches(1.9), Inches(5.9), Inches(0.9),
          "A full operations dashboard for managing warehouses, inventory, requests, and allocation decisions.",
          size=15, color=GRAY, line_spacing=1.2)
add_picture_framed(s, SS + "admin-requests.png", Inches(6.9), Inches(2.5), Inches(5.9), Inches(4.4),
                    "Admin requests screen")

# ---------- Slide 6: The Impact ----------
s = add_slide()
set_bg(s)
add_header(s, "The Impact", "What this changes for relief operations")
add_bullets(s, Inches(0.55), Inches(1.5), Inches(6.2), Inches(5.5),
    [
        ("Faster response times.", "Requests move from submission to allocation in a live pipeline instead of chat threads or paperwork."),
        ("Better resource allocation.", "Coordinators always see current inventory and allocate stock against real, verified requests."),
        ("Prioritized triage.", "Urgency levels such as critical and high surface the most urgent needs first."),
        ("Transparency for everyone.", "Volunteers track the real time status of the requests they submitted."),
        ("Data driven oversight.", "An analytics dashboard gives coordinators an at a glance view of active requests and status distribution."),
    ], size=15.5, space_after=14)
add_picture_framed(s, SS + "admin-dashboard.png", Inches(7.1), Inches(1.55), Inches(5.65), Inches(5.4),
                    "Admin analytics dashboard")

# ---------- Slide 7: The Technology ----------
s = add_slide()
set_bg(s)
add_header(s, "The Technology", "Built on Flutter and Firebase")

rows = [
    ("App framework", "Flutter (Dart), Material 3 UI"),
    ("Backend and database", "Firebase Cloud Firestore, real time streams"),
    ("Authentication", "Firebase Authentication, role based access"),
    ("File storage", "Firebase Storage for request photos"),
    ("Maps", "flutter_map with OpenStreetMap tiles, latlong2"),
    ("Analytics and charts", "fl_chart"),
    ("State management", "provider"),
    ("Device features", "geolocator for location, image_picker for photo capture"),
]
table_x, table_y = Inches(0.55), Inches(1.5)
table_w, table_h = Inches(7.6), Inches(5.3)
rows_n = len(rows) + 1
tbl_shape = s.shapes.add_table(rows_n, 2, table_x, table_y, table_w, table_h)
tbl = tbl_shape.table
tbl.columns[0].width = Inches(2.7)
tbl.columns[1].width = Inches(4.9)
hdr = ["Layer", "Technology"]
for c, text in enumerate(hdr):
    cell = tbl.cell(0, c)
    cell.text = text
    cell.fill.solid()
    cell.fill.fore_color.rgb = NAVY
    for p in cell.text_frame.paragraphs:
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = WHITE
for r, (a, b) in enumerate(rows, start=1):
    for c, text in enumerate([a, b]):
        cell = tbl.cell(r, c)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = LIGHT_GRAY if r % 2 == 0 else WHITE
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(12.5)
            p.font.color.rgb = GRAY
            p.font.bold = (c == 0)

add_text(s, Inches(8.5), Inches(1.5), Inches(4.3), Inches(0.5),
          "Architecture highlights", size=17, color=NAVY, bold=True)
add_bullets(s, Inches(8.5), Inches(2.05), Inches(4.3), Inches(4.7),
    [
        "Role based access control enforced in the UI and in Firestore security rules.",
        "Real time data sync via Firestore streams, with no manual refresh needed.",
        "Cross platform from a single Dart codebase, covering web, Android, iOS, and desktop.",
    ], size=14, space_after=16)

# ---------- Slide 8: What We Built ----------
s = add_slide()
set_bg(s)
add_header(s, "What We Actually Built", "End to end functionality across both roles")
add_bullets(s, Inches(0.55), Inches(1.5), Inches(11.9), Inches(5.5),
    [
        ("Authentication and onboarding:", "sign up with role selection, secure login, and an auth gate that routes users to the correct experience."),
        ("Field Volunteer flow:", "new request form, \"My Requests\" screen with real time status tracking, and a map view of submitted requests."),
        ("Coordinator / Admin flow:", "analytics dashboard, filterable requests screen with one tap stock allocation, inventory management, and an interactive map."),
        ("Shared infrastructure:", "data models for users, requests, inventory, and warehouses; a Firestore service layer; reusable UI components; an app wide Material 3 theme."),
        ("Demo ready:", "seeded demo accounts and sample data across Food, Water, Medicine, Shelter, and Clothing for live walkthroughs."),
    ], size=16, space_after=16)

# ---------- Slide 9: Screenshots ----------
s = add_slide()
set_bg(s)
add_header(s, "Product Walkthrough", "Screenshots")
imgs = [
    (SS + "admin-dashboard.png", "Admin dashboard"),
    (SS + "admin-map.png", "Admin map"),
    (SS + "volunteer-home.png", "New request form"),
    (SS + "volunteer-map.png", "Volunteer map"),
]
w = Inches(2.85)
h = Inches(4.9)
gap = Inches(0.25)
total_w = w * 4 + gap * 3
start_x = int((prs.slide_width - total_w) / 2)
for i, (path, cap) in enumerate(imgs):
    x = start_x + i * (w + gap)
    add_picture_framed(s, path, x, Inches(1.55), w, h, cap)

# ---------- Slide 10: Closing ----------
s = add_slide()
set_bg(s, NAVY)
add_rect(s, 0, Inches(3.9), prs.slide_width, Inches(0.06), TEAL)
add_text(s, Inches(1), Inches(2.3), Inches(11.3), Inches(1.4),
          "Current Status", size=32, color=RGBColor(0xB9, 0xD3, 0xE8), bold=True)
add_text(s, Inches(1), Inches(3.0), Inches(11.3), Inches(1.6),
          "Fully functional end to end prototype covering both user roles, running on web and "
          "mobile from a single Flutter codebase, backed by live Firebase services.",
          size=20, color=WHITE, line_spacing=1.3)
add_text(s, Inches(1), Inches(5.4), Inches(11.3), Inches(0.6),
          "Thank you. Questions welcome.", size=18, color=RGBColor(0x8C, 0xA8, 0xC2))

prs.save("Relief_Dashboard_Presentation.pptx")
print("done")
